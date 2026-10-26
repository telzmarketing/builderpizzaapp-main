import csv
import io
import uuid
from typing import Optional

from fastapi import APIRouter, Depends, File, Header, HTTPException, Query, Request, UploadFile
from sqlalchemy.exc import IntegrityError
from sqlalchemy import desc
from sqlalchemy.orm import Session

from backend.database import get_db
from backend.routes.admin_auth import get_current_admin
from backend.routes.customer_access import require_customer_or_admin
from backend.models.admin import AdminUser
from backend.models.customer import Customer, Address
from backend.models.customer_identity import CustomerChannel
from backend.models.order import Order
from backend.models.customer_event import CustomerEvent
from backend.services.customer_identity_service import (
    CustomerIdentityService,
    is_system_lead_email,
    normalize_phone,
)
from backend.services.customer_ai_service import (
    accept_customer_ai_suggestion,
    analyze_customer_profile,
    get_customer_ai_profile,
    list_customer_ai_suggestions,
    reject_customer_ai_suggestion,
)
from backend.schemas.customer import (
    CustomerCreate, CustomerUpdate, CustomerOut,
    AddressCreate, AddressOut, CustomerIdentityOut, WhatsAppLeadCreate, AdminCustomerCreate,
)
from backend.core.response import ok
from backend.core.tenant_context import TenantContext
from backend.core.tenant_ownership import assign_tenant_on_create, customers_orders_enforcement_enabled, scope_query_to_tenant
from backend.core.tenant_runtime import resolve_panel_tenant_context, resolve_public_tenant_context
from backend.core.rbac_authorization import AuthorizedTenantActor, require_rbac_permission

router = APIRouter(prefix="/customers", tags=["customers"])
suggestions_router = APIRouter(prefix="/suggestions", tags=["customer-ai"])
admin_router = APIRouter(prefix="/admin/customers", tags=["admin-customers"])

_IMPORT_HEADERS = {"name", "nome", "nome completo", "phone", "telefone", "whatsapp", "email", "e-mail"}


def _parse_customer_import(filename: str, raw: bytes) -> list[tuple[int, dict[str, str]]]:
    name = (filename or "").lower().strip()
    if len(raw) > 5 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="Arquivo de clientes excede o limite de 5 MB.")
    if name.endswith(".csv"):
        text = None
        for encoding in ("utf-8-sig", "latin-1"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise HTTPException(status_code=422, detail="CSV inválido. Use UTF-8 ou Latin-1.")
        try:
            dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        if not reader.fieldnames:
            raise HTTPException(status_code=422, detail="CSV sem cabeçalho.")
        return [(index, {str(k or "").strip().lower(): str(v or "").strip() for k, v in row.items()})
                for index, row in enumerate(reader, start=2) if any(str(v or "").strip() for v in row.values())]
    if name.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise HTTPException(status_code=503, detail="Importação XLSX indisponível no servidor.") from exc
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        rows = workbook.active.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            raise HTTPException(status_code=422, detail="Planilha sem cabeçalho.")
        normalized_headers = [str(value or "").strip().lower() for value in headers]
        return [(index, {key: str(value or "").strip() for key, value in zip(normalized_headers, values)})
                for index, values in enumerate(rows, start=2) if any(str(v or "").strip() for v in values)]
    raise HTTPException(status_code=422, detail="Envie um arquivo CSV ou XLSX.")


def _import_value(row: dict[str, str], *keys: str) -> str:
    for key in keys:
        value = row.get(key)
        if value:
            return value.strip()
    return ""

LEGACY_TENANT_ID = "tenant-legacy-default"


def _resolved_tenant_id(context: TenantContext | None) -> str:
    return context.tenant_id if context is not None else LEGACY_TENANT_ID


def _phone_conflict(db: Session, tenant_id: str, phone: str | None) -> Customer | None:
    normalized = normalize_phone(phone)
    if not normalized:
        return None
    return (
        db.query(Customer)
        .filter(Customer.tenant_id == tenant_id, Customer.normalized_phone == normalized)
        .first()
        or db.query(Customer)
        .filter(Customer.tenant_id == tenant_id, Customer.phone.in_([normalized, phone]))
        .first()
    )


def _raise_phone_conflict(customer: Customer) -> None:
    raise HTTPException(status_code=409, detail={
        "code": "CustomerPhoneConflict",
        "message": "Já existe um cliente com este telefone.",
        "customer": {
            "id": customer.id,
            "name": customer.name,
            "phone": customer.phone,
            "email": customer.email,
        },
    })


def _tenant_query(db: Session, model, context: TenantContext | None):
    return scope_query_to_tenant(
        db.query(model), model, context,
        enabled=customers_orders_enforcement_enabled(),
    )


def _own(resource, context: TenantContext | None):
    return assign_tenant_on_create(
        resource, context, enabled=customers_orders_enforcement_enabled()
    )


def _profile_level(customer: Customer) -> str:
    if customer.password_hash and not is_system_lead_email(customer.email):
        return "complete"
    if customer.phone and customer.addresses:
        return "partial"
    return "lead"


def _identity_payload(db: Session, customer: Customer, created: bool = False) -> dict:
    channel_query = db.query(CustomerChannel).filter(
        CustomerChannel.customer_id == customer.id,
        CustomerChannel.channel == "whatsapp",
    )
    if customers_orders_enforcement_enabled():
        channel_query = channel_query.filter(CustomerChannel.tenant_id == customer.tenant_id)
    channel = channel_query.first()
    return CustomerIdentityOut(
        customer=CustomerOut.model_validate(customer),
        channel=channel,
        created=created,
        profile_level=_profile_level(customer),
    ).model_dump()


@router.get("", response_model=list[CustomerOut])
def list_customers(
    request: Request,
    db: Session = Depends(get_db),
    admin: AdminUser = Depends(get_current_admin),
):
    context = resolve_panel_tenant_context(request, db, admin)
    return _tenant_query(db, Customer, context).order_by(Customer.name).all()


@router.get("/identity/by-phone")
def get_customer_identity_by_phone(
    request: Request,
    phone: str = Query(..., min_length=8),
    channel: str | None = Query(default="whatsapp"),
    db: Session = Depends(get_db),
    admin: AdminUser = Depends(get_current_admin),
):
    context = resolve_panel_tenant_context(request, db, admin)
    normalized = normalize_phone(phone)
    identity = CustomerIdentityService(db)
    tenant_id = context.tenant_id if context else None
    customer = identity.find_by_phone(normalized, channel=channel, tenant_id=tenant_id) or identity.find_by_phone(
        normalized, tenant_id=tenant_id
    )
    if not customer:
        return ok(None, "Cliente nao encontrado para este telefone.")
    return ok(_identity_payload(db, customer))


@router.post("/identity/whatsapp-lead", status_code=201)
def create_whatsapp_lead(
    body: WhatsAppLeadCreate,
    request: Request,
    db: Session = Depends(get_db),
    admin: AdminUser = Depends(get_current_admin),
):
    context = resolve_panel_tenant_context(request, db, admin)
    tenant_id = context.tenant_id if context else None
    try:
        customer, created = CustomerIdentityService(db).get_or_create_whatsapp_lead(
            phone=body.phone,
            name=body.name,
            source=body.source or "whatsapp",
            tenant_id=tenant_id,
        )
        if created and tenant_id:
            from backend.services.automation_event_producer import AutomationEventProducer
            AutomationEventProducer(db, tenant_id).customer_created(customer)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    db.commit()
    db.refresh(customer)
    return ok(
        _identity_payload(db, customer, created=created),
        "Lead WhatsApp criado." if created else "Cliente existente localizado pelo telefone.",
    )


@router.get("/{customer_id}", response_model=CustomerOut)
def get_customer(
    customer_id: str,
    request: Request,
    db: Session = Depends(get_db),
    admin: AdminUser = Depends(get_current_admin),
):
    context = resolve_panel_tenant_context(request, db, admin)
    customer = _tenant_query(db, Customer, context).filter(Customer.id == customer_id).first()
    if not customer:
        raise HTTPException(404, "Cliente não encontrado.")
    return customer


@router.post("", response_model=CustomerOut, status_code=201)
def create_customer(body: CustomerCreate, request: Request, db: Session = Depends(get_db)):
    context = resolve_public_tenant_context(request, db)
    tenant_id = _resolved_tenant_id(context)
    normalized_phone = normalize_phone(body.phone)
    existing_phone = _phone_conflict(db, tenant_id, normalized_phone)
    if existing_phone:
        _raise_phone_conflict(existing_phone)
    existing = _tenant_query(db, Customer, context).filter(Customer.email == body.email).first()
    if existing:
        raise HTTPException(400, "E-mail já cadastrado.")
    customer = Customer(id=str(uuid.uuid4()), normalized_phone=normalized_phone or None, **body.model_dump())
    _own(customer, context)
    db.add(customer)
    try:
        db.flush()
    except IntegrityError as exc:
        db.rollback()
        existing_phone = _phone_conflict(db, tenant_id, normalized_phone)
        if existing_phone:
            _raise_phone_conflict(existing_phone)
        raise exc
    CustomerIdentityService(db).sync_registered_customer(customer, auth_provider="manual")
    tenant_id = context.tenant_id if context else customer.tenant_id
    if tenant_id:
        from backend.services.automation_event_producer import AutomationEventProducer
        AutomationEventProducer(db, tenant_id).customer_created(customer)
    db.commit()
    db.refresh(customer)
    return customer


@admin_router.post("", response_model=CustomerOut, status_code=201)
def create_admin_customer(
    body: AdminCustomerCreate,
    db: Session = Depends(get_db),
    actor: AuthorizedTenantActor = Depends(require_rbac_permission("clientes", "create")),
):
    """Create a minimal customer for operator flows inside one tenant."""
    normalized_phone = normalize_phone(body.phone)
    existing_phone = _phone_conflict(db, actor.tenant_id, normalized_phone)
    if existing_phone:
        _raise_phone_conflict(existing_phone)
    email = str(body.email).strip().lower() if body.email else None
    if email and db.query(Customer).filter(Customer.tenant_id == actor.tenant_id, Customer.email == email).first():
        raise HTTPException(status_code=409, detail={
            "code": "CustomerEmailConflict",
            "message": "Já existe um cliente com este e-mail.",
        })
    customer = Customer(
        id=str(uuid.uuid4()),
        tenant_id=actor.tenant_id,
        name=body.name.strip(),
        email=email,
        phone=body.phone,
        normalized_phone=normalized_phone or None,
        source="admin",
        lgpd_consent=False,
        marketing_email_consent=False,
        marketing_whatsapp_consent=False,
    )
    db.add(customer)
    try:
        db.flush()
        if body.address is not None:
            address = Address(
                id=str(uuid.uuid4()), tenant_id=actor.tenant_id,
                customer_id=customer.id, **body.address.model_dump(),
            )
            db.add(address)
        CustomerIdentityService(db).sync_registered_customer(customer, auth_provider="none")
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        existing_phone = _phone_conflict(db, actor.tenant_id, normalized_phone)
        if existing_phone:
            _raise_phone_conflict(existing_phone)
        raise exc
    db.refresh(customer)
    return customer


@admin_router.put("/{customer_id}", response_model=CustomerOut)
def update_admin_customer(
    customer_id: str,
    body: CustomerUpdate,
    db: Session = Depends(get_db),
    actor: AuthorizedTenantActor = Depends(require_rbac_permission("clientes", "edit")),
):
    """Update customer identity fields inside the authenticated tenant."""
    customer = db.query(Customer).filter(
        Customer.id == customer_id, Customer.tenant_id == actor.tenant_id,
    ).first()
    if customer is None:
        raise HTTPException(404, "Cliente não encontrado.")
    updates = body.model_dump(exclude_none=True)
    address_payload = updates.pop("address", None)
    if "phone" in updates:
        normalized = normalize_phone(updates["phone"])
        existing = _phone_conflict(db, actor.tenant_id, normalized)
        if existing and existing.id != customer.id:
            _raise_phone_conflict(existing)
        updates["normalized_phone"] = normalized or None
    if "email" in updates and updates["email"]:
        email = str(updates["email"]).strip().lower()
        existing = db.query(Customer).filter(
            Customer.tenant_id == actor.tenant_id,
            Customer.email == email,
            Customer.id != customer.id,
        ).first()
        if existing:
            raise HTTPException(status_code=409, detail={"code": "CustomerEmailConflict", "message": "Já existe um cliente com este e-mail."})
        updates["email"] = email
    for key, value in updates.items():
        setattr(customer, key, value)
    if address_payload is not None:
        existing_address = db.query(Address).filter(
            Address.customer_id == customer.id,
            Address.tenant_id == actor.tenant_id,
            Address.is_default.is_(True),
        ).first()
        if existing_address is None:
            existing_address = db.query(Address).filter(
                Address.customer_id == customer.id,
                Address.tenant_id == actor.tenant_id,
            ).first()
        if existing_address is None:
            existing_address = Address(id=str(uuid.uuid4()), tenant_id=actor.tenant_id, customer_id=customer.id, **address_payload)
            db.add(existing_address)
        else:
            for key, value in address_payload.items():
                setattr(existing_address, key, value)
    db.commit()
    db.refresh(customer)
    return customer


@admin_router.get("/import-template")
def download_customer_import_template(
    actor: AuthorizedTenantActor = Depends(require_rbac_permission("clientes", "create")),
):
    """Download the exact headers accepted by the customer importer."""
    return {
        "filename": "modelo-clientes.csv",
        "content": "nome,telefone,email\nJoão da Silva,5511999999999,joao@example.com\n",
    }


@admin_router.post("/import")
async def import_customers(
    file: UploadFile = File(...),
    commit: bool = Query(default=False),
    db: Session = Depends(get_db),
    actor: AuthorizedTenantActor = Depends(require_rbac_permission("clientes", "create")),
):
    """Preview or commit a tenant-scoped customer import.

    Phone is the deduplication key. Existing records are never overwritten and
    imported rows never grant marketing consent.
    """
    raw = await file.read()
    rows = _parse_customer_import(file.filename or "", raw)
    existing = {
        row[0]
        for row in db.query(Customer.normalized_phone).filter(
            Customer.tenant_id == actor.tenant_id,
            Customer.normalized_phone.isnot(None),
        ).all()
        if row[0]
    }
    seen: set[str] = set()
    valid: list[dict] = []
    errors: list[dict] = []
    for row_number, row in rows:
        name = _import_value(row, "nome", "nome completo", "name")
        phone = _import_value(row, "telefone", "whatsapp", "phone")
        email = _import_value(row, "e-mail", "email") or None
        normalized = normalize_phone(phone)
        if not name:
            errors.append({"row": row_number, "reason": "Nome obrigatório."})
        elif not normalized:
            errors.append({"row": row_number, "reason": "Telefone obrigatório e inválido."})
        elif normalized in existing or normalized in seen:
            errors.append({"row": row_number, "reason": "Telefone já cadastrado ou duplicado no arquivo."})
        else:
            seen.add(normalized)
            valid.append({"name": name[:200], "phone": phone[:30], "normalized_phone": normalized, "email": email})
    created = 0
    if commit and valid:
        for payload in valid:
            customer = Customer(
                id=str(uuid.uuid4()), tenant_id=actor.tenant_id,
                name=payload["name"], phone=payload["phone"],
                normalized_phone=payload["normalized_phone"], email=payload["email"],
                source="import", lgpd_consent=False,
                marketing_email_consent=False, marketing_whatsapp_consent=False,
            )
            try:
                with db.begin_nested():
                    db.add(customer)
                    db.flush()
                created += 1
            except IntegrityError:
                errors.append({"row": None, "reason": "Telefone já cadastrado durante a importação."})
        db.commit()
    return {
        "mode": "committed" if commit else "preview",
        "total_rows": len(rows),
        "valid_count": len(valid),
        "created_count": created,
        "rejected_count": len(errors),
        "rows": valid[:100],
        "errors": errors[:500],
    }


@router.put("/{customer_id}", response_model=CustomerOut)
def update_customer(
    customer_id: str,
    body: CustomerUpdate,
    request: Request,
    db: Session = Depends(get_db),
    authorization: str | None = Header(default=None),
    x_customer_phone: str | None = Header(default=None),
    x_customer_email: str | None = Header(default=None),
):
    context = resolve_public_tenant_context(request, db)
    customer = _tenant_query(db, Customer, context).filter(Customer.id == customer_id).first()
    if not customer:
        raise HTTPException(404, "Cliente não encontrado.")
    require_customer_or_admin(
        customer, db, authorization, x_customer_phone, x_customer_email,
        expected_tenant_id=context.tenant_id if context else None,
    )
    updates = body.model_dump(exclude_none=True)
    if "phone" in updates:
        normalized_phone = normalize_phone(updates["phone"])
        existing_phone = _phone_conflict(db, _resolved_tenant_id(context), normalized_phone)
        if existing_phone and existing_phone.id != customer.id:
            _raise_phone_conflict(existing_phone)
        updates["normalized_phone"] = normalized_phone or None
    for key, value in updates.items():
        setattr(customer, key, value)
    db.commit()
    db.refresh(customer)
    return customer


# ── Addresses ─────────────────────────────────────────────────────────────────

@router.get("/{customer_id}/addresses", response_model=list[AddressOut])
def list_addresses(
    customer_id: str,
    request: Request,
    db: Session = Depends(get_db),
    authorization: str | None = Header(default=None),
    x_customer_phone: str | None = Header(default=None),
    x_customer_email: str | None = Header(default=None),
):
    context = resolve_public_tenant_context(request, db)
    customer = _tenant_query(db, Customer, context).filter(Customer.id == customer_id).first()
    if not customer:
        raise HTTPException(404, "Cliente não encontrado.")
    require_customer_or_admin(
        customer, db, authorization, x_customer_phone, x_customer_email,
        expected_tenant_id=context.tenant_id if context else None,
    )
    return _tenant_query(db, Address, context).filter(Address.customer_id == customer_id).all()


@router.post("/{customer_id}/addresses", response_model=AddressOut, status_code=201)
def add_address(
    customer_id: str,
    body: AddressCreate,
    request: Request,
    db: Session = Depends(get_db),
    authorization: str | None = Header(default=None),
    x_customer_phone: str | None = Header(default=None),
    x_customer_email: str | None = Header(default=None),
):
    context = resolve_public_tenant_context(request, db)
    customer = _tenant_query(db, Customer, context).filter(Customer.id == customer_id).first()
    if not customer:
        raise HTTPException(404, "Cliente não encontrado.")
    require_customer_or_admin(
        customer, db, authorization, x_customer_phone, x_customer_email,
        expected_tenant_id=context.tenant_id if context else None,
    )
    if body.is_default:
        _tenant_query(db, Address, context).filter(Address.customer_id == customer_id).update({"is_default": False})
    address = Address(id=str(uuid.uuid4()), customer_id=customer_id, **body.model_dump())
    _own(address, context)
    db.add(address)
    db.commit()
    db.refresh(address)
    return address


@router.delete("/{customer_id}/addresses/{address_id}", status_code=204)
def delete_address(
    customer_id: str,
    address_id: str,
    request: Request,
    db: Session = Depends(get_db),
    authorization: str | None = Header(default=None),
    x_customer_phone: str | None = Header(default=None),
    x_customer_email: str | None = Header(default=None),
):
    context = resolve_public_tenant_context(request, db)
    customer = _tenant_query(db, Customer, context).filter(Customer.id == customer_id).first()
    if not customer:
        raise HTTPException(404, "Cliente não encontrado.")
    require_customer_or_admin(
        customer, db, authorization, x_customer_phone, x_customer_email,
        expected_tenant_id=context.tenant_id if context else None,
    )
    address = _tenant_query(db, Address, context).filter(
        Address.id == address_id, Address.customer_id == customer_id
    ).first()
    if not address:
        raise HTTPException(404, "Endereço não encontrado.")
    db.delete(address)
    db.commit()


# ── Customer orders ────────────────────────────────────────────────────────────

@router.get("/{customer_id}/orders")
def get_customer_orders(
    customer_id: str,
    request: Request,
    limit: int = Query(default=50, ge=1, le=200),
    db: Session = Depends(get_db),
    authorization: str | None = Header(default=None),
    x_customer_phone: str | None = Header(default=None),
    x_customer_email: str | None = Header(default=None),
):
    context = resolve_public_tenant_context(request, db)
    customer = _tenant_query(db, Customer, context).filter(Customer.id == customer_id).first()
    if not customer:
        raise HTTPException(404, "Cliente não encontrado.")
    require_customer_or_admin(
        customer, db, authorization, x_customer_phone, x_customer_email,
        expected_tenant_id=context.tenant_id if context else None,
    )

    orders = (
        _tenant_query(db, Order, context)
        .filter(Order.customer_id == customer_id)
        .order_by(desc(Order.created_at))
        .limit(limit)
        .all()
    )

    result = []
    for o in orders:
        status_val = o.status.value if hasattr(o.status, "value") else str(o.status)
        items = []
        for item in o.items:
            flavors = [{"name": f.flavor_name, "price": f.flavor_price} for f in item.flavors] if item.flavors else []
            items.append({
                "id": item.id,
                "product_id": item.product_id,
                "quantity": item.quantity,
                "selected_size": item.selected_size,
                "selected_crust_type": item.selected_crust_type,
                "selected_drink_variant": item.selected_drink_variant,
                "notes": item.notes,
                "flavors": flavors,
                "unit_price": item.unit_price,
            })
        result.append({
            "id": o.id,
            "status": status_val,
            "total": o.total,
            "subtotal": o.subtotal,
            "shipping_fee": o.shipping_fee,
            "discount": o.discount,
            "created_at": o.created_at,
            "paid_at": o.paid_at,
            "delivered_at": o.delivered_at,
            "delivery_name": o.delivery_name,
            "delivery_street": o.delivery_street,
            "delivery_city": o.delivery_city,
            "delivery_complement": o.delivery_complement,
            "coupon_id": o.coupon_id,
            "total_time_minutes": o.total_time_minutes,
            "estimated_time": o.estimated_time,
            "items": items,
            "payment_status": o.payment.status.value if o.payment else "pending",
        })
    return ok(result)


# ── Customer events timeline ──────────────────────────────────────────────────

@router.get("/{customer_id}/events")
def get_customer_events(
    customer_id: str,
    request: Request,
    event_type: Optional[str] = Query(default=None),
    limit: int = Query(default=100, ge=1, le=500),
    db: Session = Depends(get_db),
    admin: AdminUser = Depends(get_current_admin),
):
    context = resolve_panel_tenant_context(request, db, admin)
    customer = _tenant_query(db, Customer, context).filter(Customer.id == customer_id).first()
    if not customer:
        raise HTTPException(404, "Cliente não encontrado.")

    q = _tenant_query(db, CustomerEvent, context).filter(CustomerEvent.customer_id == customer_id)
    if event_type:
        q = q.filter(CustomerEvent.event_type == event_type)

    events = q.order_by(desc(CustomerEvent.created_at)).limit(limit).all()

    return ok([{
        "id": e.id,
        "session_id": e.session_id,
        "event_type": e.event_type,
        "event_name": e.event_name,
        "event_description": e.event_description,
        "product_id": e.product_id,
        "order_id": e.order_id,
        "campaign_id": e.campaign_id,
        "coupon_id": e.coupon_id,
        "metadata_json": e.metadata_json,
        "source": e.source,
        "utm_source": e.utm_source,
        "utm_campaign": e.utm_campaign,
        "device_type": e.device_type,
        "browser": e.browser,
        "page_url": e.page_url,
        "created_at": e.created_at,
    } for e in events])


# ── Customer behavioral summary ────────────────────────────────────────────────

@router.get("/{customer_id}/summary")
def get_customer_summary(
    customer_id: str,
    request: Request,
    db: Session = Depends(get_db),
    admin: AdminUser = Depends(get_current_admin),
):
    context = resolve_panel_tenant_context(request, db, admin)
    customer = _tenant_query(db, Customer, context).filter(Customer.id == customer_id).first()
    if not customer:
        raise HTTPException(404, "Cliente não encontrado.")

    orders = _tenant_query(db, Order, context).filter(Order.customer_id == customer_id).all()
    total_orders = len(orders)
    total_spent = sum(o.total or 0 for o in orders)
    avg_ticket = total_spent / total_orders if total_orders > 0 else 0

    def _status(o: Order) -> str:
        return o.status.value if hasattr(o.status, "value") else str(o.status)

    paid_count = len([o for o in orders if _status(o) in ("paid", "pago", "preparing", "ready_for_pickup", "on_the_way", "delivered")])
    delivered_count = len([o for o in orders if _status(o) == "delivered"])
    cancelled_count = len([o for o in orders if _status(o) == "cancelled"])

    sorted_orders = sorted(orders, key=lambda o: o.created_at)
    last_order = sorted_orders[-1] if sorted_orders else None
    first_order = sorted_orders[0] if sorted_orders else None

    events = _tenant_query(db, CustomerEvent, context).filter(CustomerEvent.customer_id == customer_id).all()
    total_visits = sum(1 for e in events if e.event_type == "site_opened")
    products_viewed = sum(1 for e in events if e.event_type == "product_viewed")
    cart_abandonments = sum(1 for e in events if e.event_type == "cart_abandoned")
    checkout_abandonments = sum(1 for e in events if e.event_type == "checkout_abandoned")

    last_event = sorted(events, key=lambda e: e.created_at, reverse=True)[0] if events else None

    if total_orders >= 5:
        behavior_status = "recorrente"
    elif total_orders >= 2:
        behavior_status = "ativo"
    elif total_orders == 1:
        behavior_status = "novo_comprador"
    elif cart_abandonments > 0:
        behavior_status = "interessado"
    elif total_visits > 0:
        behavior_status = "visitante"
    else:
        behavior_status = "lead"

    return ok({
        "customer": {
            "id": customer.id,
            "name": customer.name,
            "email": customer.email,
            "phone": customer.phone,
            "crm_status": customer.crm_status,
            "source": customer.source,
            "utm_source": customer.utm_source,
            "utm_campaign": customer.utm_campaign,
            "created_at": customer.created_at,
        },
        "orders": {
            "total": total_orders,
            "total_spent": round(total_spent, 2),
            "avg_ticket": round(avg_ticket, 2),
            "paid": paid_count,
            "delivered": delivered_count,
            "cancelled": cancelled_count,
            "last_order_at": last_order.created_at if last_order else None,
            "first_order_at": first_order.created_at if first_order else None,
        },
        "behavior": {
            "status": behavior_status,
            "total_visits": total_visits,
            "products_viewed": products_viewed,
            "cart_abandonments": cart_abandonments,
            "checkout_abandonments": checkout_abandonments,
            "last_activity_at": last_event.created_at if last_event else None,
        },
    })


# AI customer profile

@router.post("/{customer_id}/analyze")
def analyze_customer(
    customer_id: str,
    request: Request,
    db: Session = Depends(get_db),
    admin: AdminUser = Depends(get_current_admin),
):
    context = resolve_panel_tenant_context(request, db, admin)
    tenant_id = context.tenant_id if context and customers_orders_enforcement_enabled() else None
    try:
        result = analyze_customer_profile(db, customer_id, tenant_id=tenant_id)
    except ValueError as exc:
        raise HTTPException(404, str(exc)) from exc
    return ok(result, "Perfil inteligente analisado.")


@router.get("/{customer_id}/profile")
def get_customer_profile(
    customer_id: str,
    request: Request,
    db: Session = Depends(get_db),
    admin: AdminUser = Depends(get_current_admin),
):
    context = resolve_panel_tenant_context(request, db, admin)
    customer = _tenant_query(db, Customer, context).filter(Customer.id == customer_id).first()
    if not customer:
        raise HTTPException(404, "Cliente nao encontrado.")
    return ok(get_customer_ai_profile(
        db, customer_id,
        tenant_id=context.tenant_id if context and customers_orders_enforcement_enabled() else None,
    ))


@router.get("/{customer_id}/suggestions")
def get_customer_suggestions(
    customer_id: str,
    request: Request,
    status: str = Query(default="pending"),
    db: Session = Depends(get_db),
    admin: AdminUser = Depends(get_current_admin),
):
    context = resolve_panel_tenant_context(request, db, admin)
    customer = _tenant_query(db, Customer, context).filter(Customer.id == customer_id).first()
    if not customer:
        raise HTTPException(404, "Cliente nao encontrado.")
    return ok(list_customer_ai_suggestions(
        db, customer_id, status=status,
        tenant_id=context.tenant_id if context and customers_orders_enforcement_enabled() else None,
    ))


@suggestions_router.post("/{suggestion_id}/accept")
def accept_suggestion(
    suggestion_id: str,
    request: Request,
    db: Session = Depends(get_db),
    admin: AdminUser = Depends(get_current_admin),
):
    context = resolve_panel_tenant_context(request, db, admin)
    admin_name = getattr(admin, "name", None) or getattr(admin, "email", None)
    try:
        result = accept_customer_ai_suggestion(
            db, suggestion_id, admin_name=admin_name,
            tenant_id=context.tenant_id if context and customers_orders_enforcement_enabled() else None,
        )
    except ValueError as exc:
        raise HTTPException(404, str(exc)) from exc
    return ok(result, "Sugestao aceita.")


@suggestions_router.post("/{suggestion_id}/reject")
def reject_suggestion(
    suggestion_id: str,
    request: Request,
    db: Session = Depends(get_db),
    admin: AdminUser = Depends(get_current_admin),
):
    context = resolve_panel_tenant_context(request, db, admin)
    try:
        result = reject_customer_ai_suggestion(
            db, suggestion_id,
            tenant_id=context.tenant_id if context and customers_orders_enforcement_enabled() else None,
        )
    except ValueError as exc:
        raise HTTPException(404, str(exc)) from exc
    return ok(result, "Sugestao rejeitada.")
